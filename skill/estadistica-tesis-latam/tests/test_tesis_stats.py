# -*- coding: utf-8 -*-
"""
Pruebas de verificación: cada cálculo de la skill se compara contra implementaciones independientes
(scipy, statsmodels, pingouin, scikit-posthocs, scikit-learn) y contra valores de referencia de G*Power.

Ejecutar:  python tests/test_tesis_stats.py     (o:  pytest tests -q)
Dependencias SOLO para probar: pingouin, scikit-posthocs, scikit-learn, pytest (opcional).
"""
import math
import os
import sys
import tempfile

import numpy as np
import pandas as pd
import pytest
from scipy import stats

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
import tesis_stats as ts  # noqa: E402

pg = pytest.importorskip("pingouin")
sp = pytest.importorskip("scikit_posthocs")
TMP = tempfile.mkdtemp()
RNG = np.random.default_rng(123)
TOL = 1e-8


def close(a, b, tol=1e-6):
    assert abs(a - b) <= tol * max(1, abs(b)), f"{a} != {b}"


# ---------------------------------------------------------------- formato
def test_formato_apa():
    assert ts.fnum(0.4567, 2, True) == ".46"
    assert ts.fnum(-0.4567, 2, True) == "-.46"
    assert ts.fnum(-0.0001, 2) == "0.00"           # sin "-0.00"
    assert ts.fp_valor(0.0004) == "< .001"
    assert ts.fp_valor(0.0234) == ".023"
    assert ts.fp_valor(1.0) == "1.000"
    assert ts.P(0.0004) == "*p* < .001" and ts.P(0.5) == "*p* = .500"
    ts.set_decimal(",")
    assert ts.fnum(1.5) == "1,50" and ts.fp_valor(0.0234) == ",023"
    ts.set_decimal(".")


def test_decision_gramatica():
    assert "se rechaza H₀" in ts.decision_txt(0.01) and "no se rechaza" not in ts.decision_txt(0.01)
    assert "no se rechaza H₀" in ts.decision_txt(0.2)
    assert "se no" not in ts.decision_txt(0.2)


# ---------------------------------------------------------------- 2 grupos
@pytest.mark.parametrize("igual", [True, False])
def test_t_independiente_vs_scipy_y_pingouin(igual):
    a, b = RNG.normal(10, 2, 31), RNG.normal(11.5, 3.5, 44)
    r = ts.t_independiente(a, b, igual)
    s = stats.ttest_ind(a, b, equal_var=igual)
    close(r["t"], s.statistic); close(r["p"], s.pvalue); close(r["gl"], s.df)
    ci = s.confidence_interval()
    close(r["ic"][0], ci.low); close(r["ic"][1], ci.high)
    close(r["d"], pg.compute_effsize(a, b, eftype="cohen"), 1e-9)


def test_mann_whitney_vs_scipy_y_pingouin():
    a = np.round(RNG.exponential(3, 40), 1)      # con empates
    b = np.round(RNG.exponential(4.5, 55), 1)
    r = ts.mann_whitney(a, b)
    s = stats.mannwhitneyu(a, b, alternative="two-sided", use_continuity=False)
    close(r["p"], s.pvalue)
    close(r["U1"], s.statistic)
    assert r["U"] == min(s.statistic, len(a) * len(b) - s.statistic)
    close(2 * stats.norm.sf(abs(r["z"])), r["p"], 1e-6)      # Z consistente con p asintótico (con empates)
    rbc = pg.mwu(a, b)["RBC"].iloc[0]
    close(abs(r["rbc"]), abs(rbc), 1e-9)


# ---------------------------------------------------------------- k grupos
def _tres():
    return [RNG.normal(m, s, n) for m, s, n in [(10, 2, 30), (12, 2.2, 28), (13, 1.8, 35)]]


def test_anova_clasica_vs_scipy_pingouin():
    g = _tres()
    r = ts.anova_clasica(g)
    s = stats.f_oneway(*g)
    close(r["F"], s.statistic); close(r["p"], s.pvalue)
    d = pd.DataFrame({"y": np.concatenate(g), "g": np.repeat(list("abc"), [len(x) for x in g])})
    a = pg.anova(d, dv="y", between="g", detailed=True)
    close(r["eta2"], a["np2"].iloc[0], 1e-9)
    close(r["ssb"], a["SS"].iloc[0], 1e-9)


def test_anova_welch_vs_pingouin_y_statsmodels():
    g = [RNG.normal(10, 1, 25), RNG.normal(12, 4, 30), RNG.normal(13, 7, 40)]
    r = ts.anova_welch(g)
    d = pd.DataFrame({"y": np.concatenate(g), "g": np.repeat(list("abc"), [len(x) for x in g])})
    w = pg.welch_anova(d, dv="y", between="g")
    close(r["F"], w["F"].iloc[0], 1e-9); close(r["gl2"], w["ddof2"].iloc[0], 1e-9); close(r["p"], w["p_unc"].iloc[0], 1e-9)
    from statsmodels.stats.oneway import anova_oneway
    sm = anova_oneway(g, use_var="unequal", welch_correction=True)
    close(r["F"], sm.statistic, 1e-9); close(r["p"], sm.pvalue, 1e-9)


def test_kruskal_y_epsilon():
    g = [np.round(RNG.exponential(m, n), 1) for m, n in [(2, 30), (3, 25), (4, 40)]]
    r = ts.kruskal(g)
    s = stats.kruskal(*g)
    close(r["H"], s.statistic); close(r["p"], s.pvalue)
    n = sum(map(len, g))
    close(r["eps2"], s.statistic / ((n ** 2 - 1) / (n + 1)))


def test_tukey_vs_pingouin():
    g = _tres()
    t = ts.tukey_hsd(g, ["a", "b", "c"])
    d = pd.DataFrame({"y": np.concatenate(g), "g": np.repeat(list("abc"), [len(x) for x in g])})
    ref = pg.pairwise_tukey(d, dv="y", between="g")
    for _, f in t.iterrows():
        A, B = f["Comparación"].split(" − ")
        row = ref[((ref["A"] == A) & (ref["B"] == B)) | ((ref["A"] == B) & (ref["B"] == A))].iloc[0]
        close(abs(f["Dif. de medias"]), abs(row["diff"]), 1e-9)
        close(f["p (Tukey)"], row["p_tukey" if "p_tukey" in row else "p-tukey"], 1e-4)


def test_games_howell_vs_pingouin():
    g = [RNG.normal(10, 1, 25), RNG.normal(12, 4, 30), RNG.normal(13, 7, 40)]
    t = ts.games_howell(g, ["a", "b", "c"])
    d = pd.DataFrame({"y": np.concatenate(g), "g": np.repeat(list("abc"), [len(x) for x in g])})
    ref = pg.pairwise_gameshowell(d, dv="y", between="g")
    for _, f in t.iterrows():
        A, B = f["Comparación"].split(" − ")
        row = ref[((ref["A"] == A) & (ref["B"] == B)) | ((ref["A"] == B) & (ref["B"] == A))].iloc[0]
        close(abs(f["Dif. de medias"]), abs(row["diff"]), 1e-9)
        pcol = [c for c in ref.columns if c.startswith("pval")][0]
        close(f["p (Games-Howell)"], row[pcol], 1e-4)


def test_dunn_bonferroni_vs_scikit_posthocs():
    g = [np.round(RNG.exponential(m, n), 1) for m, n in [(2, 30), (3, 25), (4, 40)]]
    t = ts.dunn(g, ["a", "b", "c"])
    d = pd.DataFrame({"y": np.concatenate(g), "g": np.repeat(list("abc"), [len(x) for x in g])})
    ref = sp.posthoc_dunn(d, val_col="y", group_col="g", p_adjust="bonferroni")
    for _, f in t.iterrows():
        A, B = f["Comparación"].split(" − ")
        close(f["p ajustada (Bonferroni)"], ref.loc[A, B], 1e-9)


# ---------------------------------------------------------------- pareadas
def test_wilcoxon_vs_scipy_pingouin():
    pre = np.round(RNG.exponential(3, 45), 1)
    post = np.round(pre + RNG.normal(0.6, 1.5, 45), 1)
    dif = post - pre
    r = ts.wilcoxon_pareado(dif)
    s = stats.wilcoxon(dif[dif != 0], correction=False)
    close(r["p"], s.pvalue)
    close(min(r["Wp"], r["Wm"]), s.statistic)
    ref = pg.wilcoxon(post, pre)
    close(abs(r["rbc"]), abs(ref["RBC"].iloc[0]), 1e-9)


# ---------------------------------------------------------------- correlación y chi²
def test_correlaciones_vs_scipy():
    x = RNG.normal(size=60); y = 0.5 * x + RNG.normal(size=60)
    c = ts.coef_correlacion(x, y, "pearson")
    s = stats.pearsonr(x, y)
    close(c["r"], s.statistic); close(c["p"], s.pvalue)
    ci = s.confidence_interval()
    close(c["ic"][0], ci.low, 1e-9); close(c["ic"][1], ci.high, 1e-9)
    c2 = ts.coef_correlacion(x, y, "spearman")
    s2 = stats.spearmanr(x, y)
    close(c2["r"], s2.statistic); close(c2["p"], s2.pvalue)
    assert c2["ic"][0] < c2["r"] < c2["ic"][1]


def test_chi2_vs_scipy_y_cramer():
    d = pd.DataFrame({"a": RNG.choice(list("xyz"), 200), "b": RNG.choice(["si", "no", "tal vez"], 200)})
    r = ts.chi_cuadrado(d, "a", "b", salida=TMP)
    tab = pd.crosstab(d["a"], d["b"]).values
    chi, p, gl, _ = stats.chi2_contingency(tab, correction=False)
    close(r.datos["chi2"], chi); close(r.datos["p"], p)
    close(r.datos["V"], stats.contingency.association(tab, method="cramer"), 1e-9)


def test_chi2_2x2_pocos_esperados_usa_fisher():
    d = pd.DataFrame({"a": ["u"] * 12 + ["v"] * 10, "b": ["s"] * 2 + ["n"] * 10 + ["s"] * 8 + ["n"] * 2})
    r = ts.chi_cuadrado(d, "a", "b", salida=TMP)
    assert r.datos["fisher"] is True
    tab = pd.crosstab(d["a"], d["b"]).values
    close(r.datos["p"], stats.fisher_exact(tab)[1])


# ---------------------------------------------------------------- regresión
def test_regresion_vs_statsmodels_y_betas():
    import statsmodels.formula.api as smf
    n = 120
    d = pd.DataFrame({"x1": RNG.normal(size=n), "x2": RNG.normal(5, 2, n)})
    d["y"] = 1 + 0.8 * d.x1 + 0.3 * d.x2 + RNG.normal(size=n)
    r = ts.regresion_lineal(d, "y", ["x1", "x2"], robusto=False, salida=TMP)
    m = smf.ols("y ~ x1 + x2", d).fit()
    close(r.datos["r2"], m.rsquared); close(r.datos["F"], m.fvalue); close(r.datos["p"], m.f_pvalue)
    for nom in ("x1", "x2"):
        close(r.datos["params"][nom], m.params[nom]); close(r.datos["pvalues"][nom], m.pvalues[nom])
    # β estandarizado = coeficiente de la regresión con variables z-estandarizadas
    z = d.apply(lambda c: (c - c.mean()) / c.std(ddof=1))
    mz = smf.ols("y ~ x1 + x2", z).fit()
    tb = r.datos["coef"].set_index("Variable")
    close(tb.loc["x1", "β"], mz.params["x1"], 1e-9); close(tb.loc["x2", "β"], mz.params["x2"], 1e-9)
    # VIF = 1/(1-R²) de regresar x1 sobre x2
    r2 = smf.ols("x1 ~ x2", d).fit().rsquared
    close(r.datos["vif"]["x1"], 1 / (1 - r2), 1e-8)


def test_regresion_robusta_si_heterocedasticidad():
    n = 300
    x = RNG.uniform(1, 10, n)
    d = pd.DataFrame({"x": x, "y": 2 + 0.5 * x + RNG.normal(0, 0.15 * x ** 1.6, n)})
    r = ts.regresion_lineal(d, "y", "x", salida=TMP)
    assert r.datos["bp_p"] < 0.05 and r.datos["robusto"] is True


def test_regresion_categorica_dummies_y_referencia():
    n = 90
    d = pd.DataFrame({"g": RNG.choice(["A", "B", "C"], n), "x": RNG.normal(size=n)})
    d["y"] = d["x"] + (d["g"] == "C") * 2 + RNG.normal(size=n)
    r = ts.regresion_lineal(d, "y", ["g", "x"], referencia={"g": "B"}, salida=TMP)
    nombres = list(r.datos["coef"]["Variable"])
    assert "g [A]" in nombres and "g [C]" in nombres and "g [B]" not in nombres


# ---------------------------------------------------------------- logística
def test_logistica_vs_statsmodels_y_metricas():
    import statsmodels.api as sm
    from sklearn.metrics import roc_auc_score
    n = 300
    d = pd.DataFrame({"x1": RNG.normal(size=n), "x2": RNG.normal(size=n)})
    pr = 1 / (1 + np.exp(-(-0.3 + 1.1 * d.x1 - 0.6 * d.x2)))
    d["y"] = np.where(RNG.random(n) < pr, "evento", "no")
    r = ts.regresion_logistica(d, "y", ["x1", "x2"], evento="evento", salida=TMP)
    yb = (d["y"] == "evento").astype(float)
    m = sm.Logit(yb, sm.add_constant(d[["x1", "x2"]])).fit(disp=0)
    for k in ("x1", "x2"):
        close(r.datos["params"][k], m.params[k], 1e-6)
    cs = 1 - math.exp(2 * (m.llnull - m.llf) / n)
    close(r.datos["cox_snell"], cs, 1e-9)
    close(r.datos["nagelkerke"], cs / (1 - math.exp(2 * m.llnull / n)), 1e-9)
    close(r.datos["auc"], roc_auc_score(yb, m.predict()), 1e-9)
    close(r.datos["lr"], 2 * (m.llf - m.llnull), 1e-9)


def test_hosmer_lemeshow_calibrado_no_rechaza():
    n = 2000
    p = RNG.uniform(0.05, 0.95, n)
    y = (RNG.random(n) < p).astype(float)               # probabilidades perfectamente calibradas
    chi, gl, pv = ts.hosmer_lemeshow(y, p)
    assert gl == 8 and pv > 0.01


# ---------------------------------------------------------------- tamaño de muestra / potencia
@pytest.mark.parametrize("prueba,efecto,kw,esperado", [
    ("t_independiente", 0.5, {}, 64),          # G*Power: 64 por grupo
    ("t_independiente", 0.2, {}, 394),         # 394 por grupo
    ("t_independiente", 0.8, {}, 26),          # 26 por grupo
    ("t_pareada", 0.5, {}, 34),                # 34 pares
    ("anova", 0.25, {"k": 3}, 159),            # 159 en total
    ("anova", 0.40, {"k": 3}, 66),             # 66 en total
    ("regresion", 0.15, {"predictores": 3}, 77),   # 77 casos
    ("regresion", 0.02, {"predictores": 3}, 550),  # 550 casos
    ("chi2", 0.30, {"gl": 1}, 88),             # 88 casos
])
def test_potencia_vs_gpower(prueba, efecto, kw, esperado):
    r = ts.n_potencia(prueba, efecto, **kw)
    got = r["total"] if prueba == "anova" else r["por_grupo"]
    assert got == esperado, f"{prueba} {efecto}: {got} != {esperado}"


def test_potencia_correlacion_aproximacion_de_fisher():
    # G*Power (exacto) da 84; la aproximación de Fisher da 85: se acepta ±1
    assert abs(ts.n_potencia("correlacion", 0.30)["total"] - 84) <= 1
    assert abs(ts.n_potencia("correlacion", 0.10)["total"] - 782) <= 2


def test_muestra_poblacion_cochran():
    assert ts.n_poblacion(None)["n"] == 385
    assert ts.n_poblacion(1200)["n"] == 292
    assert ts.n_poblacion(100)["n"] == 80
    assert ts.n_poblacion(500)["n"] == 218
    assert ts.n_poblacion(10000)["n"] == 370


# ---------------------------------------------------------------- selección automática del método
def _df2(a, b):
    return pd.DataFrame({"y": np.concatenate([a, b]), "g": ["A"] * len(a) + ["B"] * len(b)})


def test_ruta_student_welch_mannwhitney():
    d = _df2(RNG.normal(10, 2, 40), RNG.normal(12, 2, 40))
    assert ts.comparar_grupos(d, "y", "g", salida=TMP).prueba.startswith("t de Student")
    d = _df2(RNG.normal(10, 1, 40), RNG.normal(12, 6, 40))
    assert ts.comparar_grupos(d, "y", "g", salida=TMP).prueba.startswith("t de Welch")
    d = _df2(RNG.exponential(2, 40), RNG.exponential(4, 40))
    assert ts.comparar_grupos(d, "y", "g", salida=TMP).prueba.startswith("U de Mann")
    d = _df2(RNG.normal(10, 2, 40), RNG.normal(12, 2, 40))
    assert ts.comparar_grupos(d, "y", "g", ordinal=True, salida=TMP).prueba.startswith("U de Mann")


def test_ruta_anova_welch_kruskal_y_posthoc():
    def d3(gens):
        return pd.DataFrame({"y": np.concatenate(gens), "g": np.repeat(list("ABC"), [len(x) for x in gens])})
    r = ts.comparar_grupos(d3([RNG.normal(10, 2, 35), RNG.normal(13, 2, 35), RNG.normal(16, 2, 35)]), "y", "g", salida=TMP)
    assert r.prueba == "ANOVA de un factor" and any("Tukey" in k for k in r.tablas)
    r = ts.comparar_grupos(d3([RNG.normal(10, 1, 35), RNG.normal(13, 5, 35), RNG.normal(16, 9, 35)]), "y", "g", salida=TMP)
    assert r.prueba == "ANOVA de Welch" and any("Games-Howell" in k for k in r.tablas)
    r = ts.comparar_grupos(d3([RNG.exponential(1, 35), RNG.exponential(3, 35), RNG.exponential(6, 35)]), "y", "g", salida=TMP)
    assert r.prueba.startswith("H de Kruskal") and any("Dunn" in k for k in r.tablas)
    # sin diferencias → no hay post hoc
    r = ts.comparar_grupos(d3([RNG.normal(10, 2, 35) for _ in range(3)]), "y", "g", salida=TMP)
    if r.datos["p"] >= 0.05:
        assert not any("post hoc" in k for k in r.tablas)


def test_normalidad_ruta_shapiro_vs_lilliefors():
    assert ts.test_normalidad(RNG.normal(size=30))["prueba"] == "Shapiro-Wilk"
    assert ts.test_normalidad(RNG.normal(size=80))["prueba"].startswith("Kolmogorov")
    assert ts.test_normalidad(RNG.normal(size=80), "shapiro")["prueba"] == "Shapiro-Wilk"
    assert ts.test_normalidad(RNG.exponential(size=200))["normal"] is False
    assert ts.test_normalidad([5, 5, 5, 5])["normal"] is False     # constante


# ---------------------------------------------------------------- errores claros
def test_errores_en_espanol():
    d = _df2(RNG.normal(size=10), RNG.normal(size=10))
    with pytest.raises(ValueError, match="No encuentro la columna"):
        ts.comparar_grupos(d, "yy", "g")
    with pytest.raises(ValueError, match="¿Quisiste decir"):
        ts.comparar_grupos(d, "Y ", "gg")
    d1 = pd.DataFrame({"y": [1, 2, 3, 4, 5], "g": ["A", "A", "A", "A", "B"]})
    with pytest.raises(ValueError, match="solo 1 observación"):
        ts.comparar_grupos(d1, "y", "g")
    with pytest.raises(ValueError, match="al menos 2 grupos"):
        ts.comparar_grupos(pd.DataFrame({"y": range(10), "g": ["A"] * 10}), "y", "g")
    with pytest.raises(ValueError, match="constante"):
        ts.correlacion(pd.DataFrame({"x": range(10), "y": [1] * 10}), "x", "y")


# ---------------------------------------------------------------- carga de datos y exportaciones
def test_carga_decimales_con_coma_y_exportaciones():
    ruta = os.path.join(TMP, "es.csv")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("grupo;nota\nA;12,5\nA;13,0\nA;11,5\nB;15,5\nB;14,0\nB;16,5\n")
    d = ts.cargar_datos(ruta)
    assert pd.api.types.is_numeric_dtype(d["nota"]) and abs(d["nota"].iloc[0] - 12.5) < 1e-12
    inf = ts.Informe("Prueba", salida=os.path.join(TMP, "inf"))
    inf.agregar(ts.comparar_grupos(d, "nota", "grupo", salida=inf.salida))
    inf.agregar(ts.muestra_poblacion(500))
    out = inf.guardar()
    for k in ("md", "xlsx", "docx"):
        assert os.path.getsize(out[k]) > 2000
    import openpyxl, docx
    wb = openpyxl.load_workbook(out["xlsx"])
    assert "Resumen" in wb.sheetnames and len(wb.sheetnames) >= 4
    doc = docx.Document(out["docx"])
    assert len(doc.tables) >= 3 and len(doc.inline_shapes) >= 1


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q", "-x"]))
