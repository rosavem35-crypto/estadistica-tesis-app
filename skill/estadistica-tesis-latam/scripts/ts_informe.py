# -*- coding: utf-8 -*-
"""
ts_informe.py — Reúne los resultados y los exporta a Word (.docx, formato de tabla APA 7), Excel y Markdown.
"""
from __future__ import annotations

import os
import re
from datetime import date

import pandas as pd

from ts_base import Resultado, _carpeta, tabla_formateada

SOFTWARE_CITA = (
    "Los análisis se realizaron con Python 3 utilizando las bibliotecas SciPy (Virtanen et al., 2020), "
    "statsmodels (Seabold y Perktold, 2010), pandas y matplotlib. El nivel de significancia fue α = .05.")


class Informe:
    def __init__(self, titulo: str = "Informe de análisis estadístico", salida: str = "resultados_estadistica"):
        self.titulo = titulo
        self.salida = _carpeta(salida)
        self.resultados: list[Resultado] = []

    def agregar(self, res: Resultado) -> Resultado:
        self.resultados.append(res)
        return res

    # ------------------------------------------------------------------ Markdown
    def markdown(self) -> str:
        L = [f"# {self.titulo}", "", f"*Generado el {date.today().strftime('%d/%m/%Y')}*", ""]
        nt, nf = 1, 1
        for r in self.resultados:
            md, nt, nf = r.a_markdown(nt, nf)
            L += [md, "---", ""]
        L += ["## Software utilizado", "", SOFTWARE_CITA, ""]
        return "\n".join(L)

    def guardar_md(self, nombre: str = "informe.md") -> str:
        ruta = os.path.join(self.salida, nombre)
        with open(ruta, "w", encoding="utf-8") as f:
            f.write(self.markdown())
        return ruta

    # ------------------------------------------------------------------ Excel
    def guardar_xlsx(self, nombre: str = "tablas.xlsx") -> str:
        from openpyxl.styles import Alignment, Border, Font, Side
        ruta = os.path.join(self.salida, nombre)
        usados = set()

        def hoja(base):
            b = re.sub(r"[\[\]\*\?/\\:]", "", base)[:28] or "Hoja"
            s, i = b, 2
            while s in usados:
                s, i = f"{b[:25]}_{i}", i + 1
            usados.add(s)
            return s

        with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
            resumen = pd.DataFrame([{"Análisis": r.titulo, "Prueba": r.prueba, "Decisión": r.decision,
                                     "Redacción de resultados": re.sub(r"\*", "", r.parrafo)} for r in self.resultados])
            resumen.to_excel(xw, sheet_name="Resumen", index=False)
            ws = xw.sheets["Resumen"]
            for col, ancho in zip("ABCD", (34, 38, 44, 110)):
                ws.column_dimensions[col].width = ancho
            for fila in ws.iter_rows(min_row=2):
                for c in fila:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in self.resultados:
                for nombre_t, df in r.tablas.items():
                    nom = hoja(f"{r.id[:14]}_{nombre_t[:12]}")
                    df.to_excel(xw, sheet_name=nom, index=False)
                    w = xw.sheets[nom]
                    borde = Side(style="thin", color="000000")
                    for c in w[1]:
                        c.font = Font(bold=True)
                        c.border = Border(top=borde, bottom=borde)
                    for j, col in enumerate(df.columns, start=1):
                        letra = w.cell(row=1, column=j).column_letter
                        w.column_dimensions[letra].width = max(12, min(40, len(str(col)) + 4))
                        for i in range(2, len(df) + 2):
                            cell = w.cell(row=i, column=j)
                            if isinstance(cell.value, float):
                                cell.number_format = "0.000" if (str(col) == "p" or str(col).startswith("p ")) else "0.00"
                        w.cell(row=len(df) + 1, column=j).border = Border(bottom=borde)
        return ruta

    # ------------------------------------------------------------------ Word
    def guardar_docx(self, nombre: str = "informe.docx") -> str:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt

        doc = Document()
        est = doc.styles["Normal"]
        est.font.name = "Times New Roman"
        est.font.size = Pt(12)
        est.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for s in doc.sections:
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)

        def runs(par, texto, size=None):
            for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
                if not parte:
                    continue
                if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
                    r = par.add_run(parte[2:-2])
                    r.bold = True
                elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
                    r = par.add_run(parte[1:-1])
                    r.italic = True
                else:
                    r = par.add_run(parte)
                if size:
                    r.font.size = Pt(size)

        from docx.shared import RGBColor

        def titulo_negro(h):
            for r_ in h.runs:
                r_.font.name = "Times New Roman"
                r_.font.color.rgb = RGBColor(0, 0, 0)

        def pegar_siguiente(par):
            par.paragraph_format.keep_with_next = True
            return par

        def borde_celda(cell, top=False, bottom=False):
            tcPr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            for lado, activo in (("top", top), ("bottom", bottom)):
                if activo:
                    e = OxmlElement(f"w:{lado}")
                    e.set(qn("w:val"), "single")
                    e.set(qn("w:sz"), "8")
                    e.set(qn("w:space"), "0")
                    e.set(qn("w:color"), "000000")
                    b.append(e)
            tcPr.append(b)

        titulo_negro(doc.add_heading(self.titulo, level=1))
        p = doc.add_paragraph()
        runs(p, f"*Generado el {date.today().strftime('%d/%m/%Y')}*")
        nt, nf = 1, 1
        for r in self.resultados:
            titulo_negro(doc.add_heading(r.titulo, level=2))
            if r.prueba:
                runs(doc.add_paragraph(), f"**Prueba aplicada:** {r.prueba}")
            if r.supuestos:
                pegar_siguiente(runs(doc.add_paragraph(), "**Verificación de supuestos**") or doc.paragraphs[-1])
                for s in r.supuestos:
                    runs(doc.add_paragraph(style="List Bullet"), s, 11)
            if r.h0 or r.h1:
                runs(doc.add_paragraph(), f"**H₀:** {r.h0}")
                runs(doc.add_paragraph(), f"**H₁:** {r.h1}")
                runs(doc.add_paragraph(), f"**Decisión:** {r.decision}")
            for nombre_t, df in r.tablas.items():
                pegar_siguiente(runs(doc.add_paragraph(), f"**Tabla {nt}**") or doc.paragraphs[-1])
                pegar_siguiente(runs(doc.add_paragraph(), f"*{nombre_t}*") or doc.paragraphs[-1])
                filas = tabla_formateada(df)
                ncol = len(filas[0])
                t = doc.add_table(rows=len(filas), cols=ncol)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.autofit = False
                # anchos proporcionales al contenido (6.5 pulgadas útiles)
                pesos = [max(11, min(24, max(len(f[j]) for f in filas) + (5 if j == 0 else 2))) for j in range(ncol)]
                total = float(sum(pesos))
                tam = 10 if ncol <= 7 else 9 if ncol <= 9 else 8
                tiene_nota = bool(r.notas_tabla.get(nombre_t))
                for j in range(ncol):
                    t.columns[j].width = Inches(6.5 * pesos[j] / total)
                for i, fila in enumerate(filas):
                    for j, val in enumerate(fila):
                        c = t.cell(i, j)
                        c.width = Inches(6.5 * pesos[j] / total)
                        c.text = ""
                        par = c.paragraphs[0]
                        par.paragraph_format.space_after = Pt(1)
                        par.paragraph_format.space_before = Pt(1)
                        par.paragraph_format.line_spacing = 1.0
                        par.paragraph_format.keep_with_next = (i < len(filas) - 1) or tiene_nota
                        run = par.add_run(val)
                        run.font.size = Pt(tam)
                        run.bold = False
                        par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                        borde_celda(c, top=(i == 0), bottom=(i == 0 or i == len(filas) - 1))
                nota = r.notas_tabla.get(nombre_t)
                if nota:
                    p = doc.add_paragraph()
                    m = re.match(r"^(Nota\.)\s*(.*)$", nota)
                    if m:
                        ri = p.add_run(m.group(1) + " ")
                        ri.italic = True
                        ri.font.size = Pt(10)
                        rr = p.add_run(m.group(2))
                        rr.font.size = Pt(10)
                    else:
                        runs(p, nota, 10)
                doc.add_paragraph()
                nt += 1
            for ruta, pie in zip(r.figuras, r.pies_figura or [""] * len(r.figuras)):
                pegar_siguiente(runs(doc.add_paragraph(), f"**Figura {nf}**") or doc.paragraphs[-1])
                pegar_siguiente(runs(doc.add_paragraph(), f"*{pie}*") or doc.paragraphs[-1])
                doc.add_picture(ruta, width=Inches(5.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                nf += 1
            if r.parrafo:
                pegar_siguiente(runs(doc.add_paragraph(), "**Redacción de resultados**") or doc.paragraphs[-1])
                runs(doc.add_paragraph(), r.parrafo)
            if r.metodo:
                pegar_siguiente(runs(doc.add_paragraph(), "**Justificación para el capítulo de método**") or doc.paragraphs[-1])
                runs(doc.add_paragraph(), r.metodo)
            if r.avisos:
                pegar_siguiente(runs(doc.add_paragraph(), "**Advertencias para el investigador**") or doc.paragraphs[-1])
                for a in r.avisos:
                    runs(doc.add_paragraph(style="List Bullet"), a, 11)
        titulo_negro(doc.add_heading("Software utilizado", level=2))
        runs(doc.add_paragraph(), SOFTWARE_CITA)
        ruta = os.path.join(self.salida, nombre)
        doc.save(ruta)
        return ruta

    def guardar(self, formatos=("md", "xlsx", "docx")) -> dict:
        out = {}
        if "md" in formatos:
            out["md"] = self.guardar_md()
        if "xlsx" in formatos:
            out["xlsx"] = self.guardar_xlsx()
        if "docx" in formatos:
            out["docx"] = self.guardar_docx()
        return out
